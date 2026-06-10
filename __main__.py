from openpyxl import load_workbook, Workbook
from openpyxl.utils.exceptions import InvalidFileException
from funcoes import menu
from rich import print
from time import sleep

def preencher_produtos(arquivo):
    produtos = list()

    for linha in arquivo.iter_rows(min_row=2, values_only=True):
        if linha[1] is not None and linha[2] is not None:
            produtos.append(
                {
                    'id': linha[0],
                    'nome': linha[1],
                    'valor': linha[2]
                }
            )
    
    return produtos


try:
    print('[yellow]Cole a planilha em exel no formato [bold blue].xlsx [yellow]nesta pasta e coloque seu nome abaixo.[/]')
    np = str(input('Digite o nome aqui: ').strip())
    if '.xlsx' not in np:
        np += '.xlsx'
    print(np)
    arq = load_workbook(np)
    planilha = arq.active

except (FileNotFoundError, InvalidFileException): 
    print('[bold red]<Pasta não encontrada, tente novamente>[/]')

else:
    print('[bold green]PASTA ENCONTRADA.[/]')
    while True:
            opcao = menu("Manipular valores", "Remover duplicatas e produtos com valores nulos", "Nada")

            if opcao == 1:
                opcao = menu("Mostrar estatísticas", "Ordenar do menor para o maior", "Voltar")
                if opcao == 1:
                    produtos = preencher_produtos(planilha)
                    
                    soma = sum(d['valor'] for d in produtos)
                    media = (soma / len(produtos))
                    print('[bold green]ESTATÍSTICAS DA PLANILHA[/]'.center(60, '-'))
                    print(f'Quantidade de produtos: {len(produtos)}\nSoma: R${soma:,.2f}\nMédia: R${media:,.2f}\nMaior: R${max(produtos, key=lambda d: d['valor'])['valor']}\nMenor: R${min(produtos, key=lambda d: d['valor'])['valor']}')

                    print('[bold green]PRODUTOS ACIMA DA MÉDIA[/]'.center(100, '-'))

                    for i, produto in enumerate(produtos):
                        if produto['valor'] > media:
                            print(f'{i+1}- <ID> {produto['id']:<5}| <[bold purple]Nome[/]> {produto['nome']:<10}| <[bold purple]Valor[/]> R${produto['valor']:<30,.2f}')

                elif opcao == 2:
                    produtos = preencher_produtos(planilha)

                    lista_ordenada = sorted(produtos, key=lambda d: d['valor'])

                    arq3 = Workbook()
                    planilha3 = arq3.active

                    linha_exel = 2
                    index = 1

                    for produto in lista_ordenada:
                        planilha3.cell(row=linha_exel, column=1, value=index)
                        planilha3.cell(row=linha_exel, column=2, value=produto['nome'])
                        planilha3.cell(row=linha_exel, column=3, value=produto['valor'])

                        linha_exel+=1
                        index+=1

                    try:
                        print('[bold yellow]Tentando salvar planilha...[/]')
                        arq3.save('produtos_ordenados.xlsx')
                    except PermissionError:
                        print('[bold red]Erro. Feche a aba do exel e execute novamente.[/]')
                    
                    else:
                        print('[bold green]PLANILHA CRIADA.[/]')
                
                elif opcao == 3:
                    pass

            elif opcao == 2:
                produtos = list()
                vistos = set()

                for linha in planilha.iter_rows(min_row=2, values_only=True):
                    try:
                        valores = (linha[1], linha[2])

                        if valores not in vistos:
                            vistos.add(valores)

                            if linha[1] is not None and linha[2] is not None:
                                if linha[2] > 0:
                                    produtos.append(
                                        {
                                            'id': linha[0],
                                            'nome': linha[1],
                                            'valor': linha[2]
                                        }
                                    )
                    except IndexError:
                        print('[bold red]Erro. Falha ao tentar utilizar uma chave.[/]')
                
                arq2 = Workbook()
                planilha2 = arq2.active

                linha_exel = 2
                indice = 1
                for produto in produtos:
                    planilha2.cell(row=linha_exel, column=1, value=indice)
                    planilha2.cell(row=linha_exel, column=2, value=produto['nome'])
                    planilha2.cell(row=linha_exel, column=3, value=produto['valor'])

                    linha_exel+=1
                    indice+=1
                
                try:
                    print('Tentando salvar nova planilha...')
                    sleep(1)
                    arq2.save("produtos_new.xlsx")

                except PermissionError:
                    print('[bold red]Erro. Feche a aba do exel e execute novamente.[/]')
                
                else:
                    print('[bold yellow]Duplicatas e valores Nulos Eliminados.[/]')
                    print('( Resultado salvo em um nova planilha para a seguranca de seus dados )')

            elif opcao == 3:
                print('Finalizando programa...')
                sleep(1.5)
                print('Até logo!')
                sleep(0.5)
                break